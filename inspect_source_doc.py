import docx

doc = docx.Document(r"C:\Users\91600\Downloads\Copy of 23IT723_Final_Project_Sample_Report.docx")

print("=== TABLES IN SOURCE DOC ===")
for t_idx, table in enumerate(doc.tables):
    rows = len(table.rows)
    cols = len(table.columns)
    style = table.style.name if table.style else "None"
    first_cell = table.rows[0].cells[0].text.strip()[:30] if rows > 0 and cols > 0 else ""
    print(f"Table {t_idx+1}: {rows}x{cols}, Style='{style}', First Cell='{first_cell}'")
    # Check cell text size
    if rows > 0:
        c = table.rows[0].cells[0]
        if c.paragraphs and c.paragraphs[0].runs:
            r = c.paragraphs[0].runs[0]
            print(f"   Header font: {r.font.name}, size={r.font.size.pt if r.font.size else None}, bold={r.font.bold}")
        if rows > 1:
            c2 = table.rows[1].cells[0]
            if c2.paragraphs and c2.paragraphs[0].runs:
                r2 = c2.paragraphs[0].runs[0]
                print(f"   Body font: {r2.font.name}, size={r2.font.size.pt if r2.font.size else None}, bold={r2.font.bold}")

print("\n=== PARAGRAPH SAMPLES ===")
samples = [
    ("Title", 5),
    ("Bonafide", 30),
    ("Heading Intro", 50),
    ("Body Intro", 52),
    ("Module Heading", 100),
    ("Figure Caption", 150)
]
for idx in range(30, min(120, len(doc.paragraphs))):
    p = doc.paragraphs[idx]
    txt = p.text.strip()
    if txt and ("INTRODUCTION" in txt or "OBJECTIVES" in txt or "ABSTRACT" in txt or "Figure" in txt or "MODULE" in txt or "TECHNOLOGY" in txt):
        r = p.runs[0] if p.runs else None
        sz = r.font.size.pt if r and r.font.size else None
        print(f"P{idx}: '{txt[:40]}' -> Font={r.font.name if r else None}, Size={sz}, Bold={r.font.bold if r else None}")
