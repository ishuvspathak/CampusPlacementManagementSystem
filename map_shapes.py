import docx

doc = docx.Document(r"C:\Users\91600\Downloads\Copy of 23IT723_Final_Project_Sample_Report.docx")

shape_idx = 0
for p_idx, p in enumerate(doc.paragraphs):
    # Check if paragraph contains an inline shape
    if 'graphic' in p._p.xml:
        shape_idx += 1
        prev_text = doc.paragraphs[p_idx-1].text.strip() if p_idx > 0 else ""
        curr_text = p.text.strip()
        next_text = doc.paragraphs[p_idx+1].text.strip() if p_idx + 1 < len(doc.paragraphs) else ""
        print(f"Shape {shape_idx} at P{p_idx}:")
        print(f"   Prev: '{prev_text[:40]}'")
        print(f"   Next: '{next_text[:40]}'")
