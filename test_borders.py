import os
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

src_path = r"C:\Users\91600\Downloads\Copy of 23IT723_Final_Project_Sample_Report.docx"
doc = docx.Document(src_path)

# Ensure all tables have explicit solid black borders (Table Grid)
def enforce_table_borders(table):
    tblPr = table._tbl.tblPr
    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'</w:tblBorders>'
    )
    # Remove existing tblBorders if present
    existing_borders = tblPr.find(docx.oxml.ns.qn('w:tblBorders'))
    if existing_borders is not None:
        tblPr.remove(existing_borders)
    tblPr.append(borders_xml)

for t in doc.tables:
    enforce_table_borders(t)

print(f"Enforced solid black table borders across all {len(doc.tables)} tables.")
